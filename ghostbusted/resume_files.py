# Copyright 2026 Asaavi Anil Tupsounder
# SPDX-License-Identifier: Apache-2.0
"""Read resume files locally (free: nothing is sent to any API).

Supported: PDF, Word (.docx), plain text (.txt / .md).
"""
from __future__ import annotations

import io
import re

MAX_BYTES = 5 * 1024 * 1024
SUPPORTED = (".pdf", ".docx", ".txt", ".md")


class ResumeFileError(ValueError):
    pass


def _clean(text: str) -> str:
    text = text.replace(" ", " ").replace("\r", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if len(data) > MAX_BYTES:
        raise ResumeFileError("That file is over 5 MB. Export a smaller PDF or Word file.")
    if name.endswith((".png", ".jpg", ".jpeg", ".webp", ".heic")):
        raise ResumeFileError("Image resumes can't be read by most applicant-tracking systems, and image upload "
                              "isn't supported yet. Upload a PDF or Word (.docx) version instead.")
    if name.endswith(".doc"):
        raise ResumeFileError("Old .doc files aren't supported. In Word, use File → Save As → .docx (or PDF).")
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
        except Exception:
            raise ResumeFileError("Couldn't read that PDF. Try exporting it again, or upload a Word file.")
    elif name.endswith(".docx"):
        import docx

        try:
            d = docx.Document(io.BytesIO(data))
        except Exception:
            raise ResumeFileError("Couldn't read that Word file. Try saving it again as .docx or PDF.")
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:  # many resume templates put content in tables
            for row in table.rows:
                parts.append("  ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ResumeFileError("Upload a PDF, Word (.docx) or text file.")
    text = _clean(text)
    if len(text) < 200:
        raise ResumeFileError("Almost no text found. If this is a scanned PDF (a picture of a resume), "
                              "export a text-based PDF or Word file from your resume editor instead.")
    return text
